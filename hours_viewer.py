"""
hours_viewer.py
===============
Entry point: collega HoursUI (widget) e hours_logic (dati).

Requisiti:
    pip install pywin32 pandas openpyxl tkcalendar Pillow

Utilizzo:
    python hours_viewer.py
    python hours_viewer.py --xlsx "\\\\server\\path\\hours_data.xlsx"
"""

import argparse
import os
import threading
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from datetime import date, datetime, timedelta
from pathlib import Path

from hours_ui import (
    HoursUI, ACCENT, ACCENT_MID, ACCENT_LIGHT,
    BG_CARD, TEXT, TEXT_DIM, TEXT_WHITE,
    SUCCESS, BORDER, ERR, WARN
)
from hours_logic import (
    XLSX_FILE, JSON_FILE,
    load_settings, save_settings,
    json_to_df, df_to_json,
    full_refresh, check_hours, parse_date,
)

import pandas as pd


# ── Icon helper ───────────────────────────────────────────────────────────────

def set_app_icon(window, icon_name: str = "HV.ico"):
    """Imposta l'icona su finestra principale e taskbar."""
    import sys
    candidates = []
    if hasattr(sys, "_MEIPASS"):
        candidates.append(os.path.join(sys._MEIPASS, icon_name))
    base = os.path.dirname(sys.executable if getattr(sys, "frozen", False)
                           else os.path.abspath(__file__))
    candidates.append(os.path.join(base, icon_name))

    for path in candidates:
        if os.path.exists(path):
            # 1. Icona barra del titolo
            try:
                window.iconbitmap(path)
            except Exception:
                pass
            # 2. AppUserModelID per taskbar Windows
            try:
                import ctypes
                ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID(
                    "Flowserve.HoursViewer.1.0")
            except Exception:
                pass
            # 3. iconphoto per taskbar (usa PIL)
            try:
                from PIL import Image, ImageTk
                img   = Image.open(path)
                photo = ImageTk.PhotoImage(img)
                window.iconphoto(True, photo)
                window._icon_photo = photo  # evita garbage collection
            except Exception:
                pass
            return


# ── Main controller ───────────────────────────────────────────────────────────

class HoursViewer:
    def __init__(self, xlsx_path: str):
        self.xlsx_path     = xlsx_path
        self.df_full       = None
        self._last_df      = None
        self._hours_issues = []

        self.ui = HoursUI(None)
        set_app_icon(self.ui.root)
        self._wire_callbacks()
        # Salva il percorso xlsx nel settings
        s = load_settings()
        s["xlsx_path"] = xlsx_path
        save_settings(s)
        self._trigger_load()
        self.ui.lbl_footer.config(text=f"Cache: {JSON_FILE}  │  Data: {xlsx_path}")

    # ── Wire callbacks ────────────────────────────────────────────────────────

    def _wire_callbacks(self):
        ui = self.ui
        ui.btn_reload.config(command=self._trigger_refresh)
        ui.btn_cerca.config(command=self._search)
        ui._btn_reset.config(command=self._reset)
        ui._btn_export.config(command=self._export)
        ui.btn_detail.config(command=self._show_hours_detail)
        ui.btn_cal_from.config(command=lambda: self._pick_date(ui.ent_from))
        ui.btn_cal_to.config(command=lambda: self._pick_date(ui.ent_to))
        ui.cmb_gruppo.bind("<<ComboboxSelected>>", self._on_gruppo_change)
        for col in ("Resource", "Job", "Phase", "Description",
                    "Notes", "Date", "Created On", "Hours"):
            ui.tree.heading(col, command=lambda c=col: self._sort(c))
        for w in ui.frm_banner.winfo_children():
            try:
                if w.cget("text") == "Update now":
                    w.config(command=self._trigger_refresh)
            except Exception:
                pass

    # ── Load ─────────────────────────────────────────────────────────────────

    def _trigger_load(self):
        self.ui.lbl_status.config(text="● Loading...", fg=TEXT_WHITE)
        threading.Thread(target=self._load_thread, daemon=True).start()

    def _load_thread(self):
        try:
            if not JSON_FILE.exists():
                self.ui.root.after(0, lambda: self.ui.lbl_status.config(
                    text="● First run: downloading from DB...", fg=TEXT_WHITE))
                self._do_refresh()
                return
            df, exported_at = json_to_df(JSON_FILE)
            self.ui.root.after(0, lambda: self._on_ok(df, exported_at))
        except Exception as e:
            err = str(e)
            self.ui.root.after(0, lambda: self._on_err(err))

    # ── Refresh ───────────────────────────────────────────────────────────────

    def _trigger_refresh(self):
        self.ui.btn_reload.config(state="disabled")
        self.ui.btn_cerca.config(state="disabled")
        threading.Thread(target=self._do_refresh, daemon=True).start()

    def _do_refresh(self):
        def status(msg):
            self.ui.root.after(0, lambda: self.ui.lbl_status.config(
                text=f"● {msg}", fg=TEXT_WHITE))
        try:
            df = full_refresh(self.xlsx_path, status_cb=status)
            status("Saving local data...")
            df_to_json(df, JSON_FILE)
            df_saved, exported_at = json_to_df(JSON_FILE)
            self.ui.root.after(0, lambda: self._on_ok(df_saved, exported_at))
        except Exception as e:
            err = str(e)
            self.ui.root.after(0, lambda: self._on_err(err))

    # ── Callbacks ─────────────────────────────────────────────────────────────

    def _on_ok(self, df, exported_at):
        self.df_full  = df
        self._last_df = df
        self._init_controls(df)
        self.ui.lbl_status.config(text=f"● {len(df):,} records", fg=TEXT_WHITE)
        self.ui.lbl_update.config(text=f"Data as of: {exported_at}")
        self.ui.lbl_risorsa_hdr.config(text="· All Groups")
        self._restore_settings()
        self._populate(df)
        self._check_update_available()

    def _on_err(self, msg):
        self.ui.btn_reload.config(state="normal")
        self.ui.lbl_status.config(text="● Error", fg=TEXT_WHITE)
        messagebox.showerror("Error", msg)

    # ── Controls init ─────────────────────────────────────────────────────────

    def _init_controls(self, df):
        due_mesi_fa = date.today() - timedelta(days=60)
        df["CreatedOn"] = pd.to_datetime(df["CreatedOn"], errors="coerce").dt.date
        df_active = df[df["CreatedOn"] >= due_mesi_fa]

        gruppi = sorted(df_active["GroupRes"].dropna().unique().tolist()) \
                 if "GroupRes" in df.columns else []
        self.ui.cmb_gruppo["values"] = ["All"] + gruppi
        self.ui.cmb_gruppo.set("All")
        self.ui.cmb_gruppo.config(state="readonly")

        names = sorted(df_active["DescrResource"].dropna().unique().tolist())
        self.ui.cmb_name["values"] = ["All"] + names
        self.ui.cmb_name.set("All")
        self.ui.cmb_name.config(state="readonly")

        self.ui.btn_cerca.config(state="normal")
        self.ui.btn_reload.config(state="normal")

        today     = date.today()
        jan_first = date(today.year, 1, 1)
        self.ui.ent_from.delete(0, tk.END)
        self.ui.ent_from.insert(0, str(jan_first))
        self.ui.ent_to.delete(0, tk.END)
        self.ui.ent_to.insert(0, str(today))

    # ── Settings ──────────────────────────────────────────────────────────────

    def _restore_settings(self):
        s = load_settings()
        if not s:
            return
        if s.get("last_group") and s["last_group"] in self.ui.cmb_gruppo["values"]:
            self.ui.cmb_gruppo.set(s["last_group"])
            self._on_gruppo_change()
        if s.get("last_resource") and s["last_resource"] in self.ui.cmb_name["values"]:
            self.ui.cmb_name.set(s["last_resource"])
        if s.get("date_from"):
            self.ui.ent_from.delete(0, tk.END)
            self.ui.ent_from.insert(0, s["date_from"])
        if s.get("date_to"):
            self.ui.ent_to.delete(0, tk.END)
            self.ui.ent_to.insert(0, s["date_to"])

    def _save_settings(self):
        s = load_settings()
        s["last_group"]    = self.ui.cmb_gruppo.get()
        s["last_resource"] = self.ui.cmb_name.get()
        s["date_from"]     = self.ui.ent_from.get().strip()
        s["date_to"]       = self.ui.ent_to.get().strip()
        save_settings(s)

    # ── Group change ──────────────────────────────────────────────────────────

    def _on_gruppo_change(self, _event=None):
        if self.df_full is None:
            return
        due_mesi_fa = date.today() - timedelta(days=60)
        df = self.df_full.copy()
        df["CreatedOn"] = pd.to_datetime(df["CreatedOn"], errors="coerce").dt.date
        gruppo = self.ui.cmb_gruppo.get()
        if gruppo and gruppo != "All" and "GroupRes" in df.columns:
            df = df[df["GroupRes"] == gruppo]
        df_active = df[df["CreatedOn"] >= due_mesi_fa]
        names = sorted(df_active["DescrResource"].dropna().unique().tolist())
        self.ui.cmb_name["values"] = ["All"] + names
        self.ui.cmb_name.set("All")

    # ── Search / Reset ────────────────────────────────────────────────────────

    def _search(self):
        if self.df_full is None:
            return
        df = self.df_full.copy()

        gruppo = self.ui.cmb_gruppo.get()
        if gruppo and gruppo != "All" and "GroupRes" in df.columns:
            df = df[df["GroupRes"] == gruppo]

        nome = self.ui.cmb_name.get()
        if nome and nome != "All":
            df = df[df["DescrResource"] == nome]

        cdcom = self.ui.ent_cdcom.get().strip()
        if cdcom:
            df = df[df["CDCOM"].str.contains(cdcom, case=False, na=False)]

        d_from = parse_date(self.ui.ent_from.get().strip())
        d_to   = parse_date(self.ui.ent_to.get().strip())
        if d_from:
            df = df[df["WorkDte"] >= d_from]
        if d_to:
            df = df[df["WorkDte"] <= d_to]

        if df.empty:
            messagebox.showinfo("No results", "No records found.")

        if nome and nome != "All":
            grp = ""
            if "GroupRes" in self.df_full.columns:
                vals = self.df_full.loc[
                    self.df_full["DescrResource"] == nome, "GroupRes"].dropna()
                grp = vals.iloc[0] if not vals.empty else ""
            label = f"· {nome}" + (f"  |  {grp}" if grp else "")
        elif gruppo and gruppo != "All":
            label = f"· {gruppo}"
        else:
            label = "· All Groups"

        self.ui.lbl_risorsa_hdr.config(text=label)
        self._populate(df)
        self._run_hours_check()
        self._save_settings()

    def _reset(self):
        if self.df_full is None:
            return
        self.ui.cmb_gruppo.set("All")
        self._on_gruppo_change()
        self.ui.cmb_name.set("All")
        self.ui.ent_cdcom.delete(0, tk.END)
        today     = date.today()
        jan_first = date(today.year, 1, 1)
        self.ui.ent_from.delete(0, tk.END)
        self.ui.ent_from.insert(0, str(jan_first))
        self.ui.ent_to.delete(0, tk.END)
        self.ui.ent_to.insert(0, str(today))
        self.ui.lbl_risorsa_hdr.config(text="· All Groups")
        self.ui.lbl_check.config(
            text="Select a resource and period to check hours", fg=TEXT_DIM)
        self.ui.btn_detail.pack_forget()
        self._populate(self.df_full)

    # ── Hours check ───────────────────────────────────────────────────────────

    def _run_hours_check(self):
        nome   = self.ui.cmb_name.get()
        d_from = parse_date(self.ui.ent_from.get().strip())
        d_to   = parse_date(self.ui.ent_to.get().strip())

        if not nome or nome == "All" or not d_from or not d_to:
            self.ui.lbl_check.config(
                text="Select a resource and period to check hours",
                fg=TEXT_DIM)
            self.ui.btn_detail.pack_forget()
            return

        s        = load_settings()
        excluded = set(s.get("exclusions", {}).get(nome, []))
        issues   = check_hours(self.df_full, nome, d_from, d_to)
        self._hours_issues = issues

        if not issues:
            self.ui.lbl_check.config(
                text="✓  All hours complete for the selected period",
                fg=SUCCESS)
            self.ui.btn_detail.pack_forget()
        else:
            active = [i for i in issues if str(i["data"]) not in excluded]
            tot    = sum(i["delta"] for i in active)
            self.ui.lbl_check.config(
                text=f"⚠  {len(active)} days with missing hours  ({tot:.1f}h total)",
                fg=ERR)
            self.ui.btn_detail.pack(side="left")

    # ── Hours detail ──────────────────────────────────────────────────────────

    def _show_hours_detail(self):
        if not self._hours_issues:
            return

        nome     = self.ui.cmb_name.get()
        s        = load_settings()
        excluded = set(s.get("exclusions", {}).get(nome, []))

        top = tk.Toplevel(self.ui.root)
        top.geometry("500x460")
        top.configure(bg=BG_CARD)
        top.resizable(False, True)
        set_app_icon(top)

        def calc_total():
            return sum(i["delta"] for i in self._hours_issues
                       if str(i["data"]) not in excluded)

        def update_ui():
            tot = calc_total()
            lbl_tot.config(text=f"Total missing: {tot:.1f}h")
            top.title(f"Missing Hours — {nome}  |  {tot:.1f}h")
            self._run_hours_check()

        hdr = tk.Frame(top, bg=ACCENT, height=42)
        hdr.pack(fill="x")
        hdr.pack_propagate(False)
        tk.Label(hdr, text=f"Missing Hours — {nome}",
                 bg=ACCENT, fg=TEXT_WHITE,
                 font=("Segoe UI", 11, "bold")).pack(side="left", padx=16, pady=10)

        info = tk.Frame(top, bg="#fef9c3", height=26)
        info.pack(fill="x")
        info.pack_propagate(False)
        tk.Label(info,
                 text="  Click the Exclude column to exclude a day (e.g. holiday, leave)",
                 bg="#fef9c3", fg=WARN,
                 font=("Segoe UI", 8)).pack(side="left", fill="y")

        tbl_frame = tk.Frame(top, bg=BG_CARD)
        tbl_frame.pack(fill="both", expand=True, padx=12, pady=(8, 0))

        cols   = ("Exclude", "Date", "Day", "Expected", "Logged", "Missing")
        tree   = ttk.Treeview(tbl_frame, columns=cols, show="headings",
                              style="PT.Treeview", height=14)
        widths = {"Exclude": 65, "Date": 95, "Day": 55, "Expected": 80, "Logged": 80, "Missing": 75}
        for c in cols:
            tree.heading(c, text=c)
            tree.column(c, width=widths[c], anchor="center")

        tree.tag_configure("missing",  background="#fee2e2", foreground=ERR)
        tree.tag_configure("partial",  background="#fff3e0", foreground=WARN)
        tree.tag_configure("excluded", background="#f3f4f6", foreground="#9ca3af")

        for issue in sorted(self._hours_issues, key=lambda x: x["data"]):
            key  = str(issue["data"])
            excl = key in excluded
            tag  = "excluded" if excl else (
                   "missing"  if issue["ore_inserite"] == 0 else "partial")
            tree.insert("", "end", iid=key, values=(
                "☑" if excl else "☐", key, issue["giorno"],
                f"{issue['ore_attese']:.0f}h",
                f"{issue['ore_inserite']:.1f}h",
                f"-{issue['delta']:.1f}h",
            ), tags=(tag,))

        def on_click(event):
            row = tree.identify_row(event.y)
            col = tree.identify_column(event.x)
            if not row or col != "#1":
                return
            vals = list(tree.item(row)["values"])
            if row in excluded:
                excluded.discard(row)
                orig   = "missing" if vals[4] == "0.0h" else "partial"
                vals[0] = "☐"
                tree.item(row, values=vals, tags=(orig,))
            else:
                excluded.add(row)
                vals[0] = "☑"
                tree.item(row, values=vals, tags=("excluded",))
            s2 = load_settings()
            if excluded:
                s2.setdefault("exclusions", {})[nome] = list(excluded)
            else:
                # Rimuovi la voce se non ci sono esclusioni
                s2.get("exclusions", {}).pop(nome, None)
            save_settings(s2)
            update_ui()

        tree.bind("<Button-1>", on_click)

        vsb = ttk.Scrollbar(tbl_frame, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=vsb.set)
        tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        tbl_frame.rowconfigure(0, weight=1)
        tbl_frame.columnconfigure(0, weight=1)

        tk.Frame(top, bg=BORDER, height=1).pack(fill="x")
        footer = tk.Frame(top, bg=BG_CARD, height=34)
        footer.pack(fill="x")
        footer.pack_propagate(False)
        lbl_tot = tk.Label(footer, text=f"Total missing: {calc_total():.1f}h",
                           font=("Segoe UI", 10, "bold"), bg=BG_CARD, fg=ERR)
        lbl_tot.pack(side="left", padx=16, fill="y")
        top.title(f"Missing Hours — {nome}  |  {calc_total():.1f}h")

    # ── Table ─────────────────────────────────────────────────────────────────

    def _populate(self, df):
        self.ui.tree.delete(*self.ui.tree.get_children())
        self._last_df = df
        if df.empty:
            self.ui.lbl_summary.config(text="No results.")
            return

        df_s  = df.sort_values("WorkDte", ascending=False)
        total = df_s["WorkHrs"].sum()

        for i, (_, r) in enumerate(df_s.iterrows()):
            self.ui.tree.insert("", "end", values=(
                r.get("DescrResource", "") or "",
                r.get("CDCOM", "")         or "-",
                r.get("DescrFase", "")     or "-",
                r.get("DescrShort", "")    or "-",
                r.get("WorkNote", "")      or "",
                str(r.get("WorkDte", "")),
                str(r.get("CreatedOn", "")),
                f"{r.get('WorkHrs', 0):.1f}",
            ), tags=("even" if i % 2 == 0 else "odd",))

        self.ui.tree.insert("", "end", values=(
            "", "", "─" * 12, "─" * 10, "TOTAL HOURS →", "", "", f"{total:.1f}"
        ), tags=("total",))

        # Calcola ore lavorative attese nel periodo
        # (esclusi weekend, festività italiane e giorni esclusi dall'utente)
        # Usa le date dei filtri From/To, non quelle dei record
        d_from = parse_date(self.ui.ent_from.get().strip()) or df_s["WorkDte"].min()
        d_to   = parse_date(self.ui.ent_to.get().strip())   or df_s["WorkDte"].max()
        from hours_logic import is_holiday, HOURS_EXPECTED
        from datetime import timedelta

        # Leggi esclusioni per la risorsa selezionata
        nome = self.ui.cmb_name.get()
        s = load_settings()
        excluded = set(s.get("exclusions", {}).get(nome, []))

        ore_attese = 0
        current = d_from
        while current <= d_to:
            if not is_holiday(current) and str(current) not in excluded:
                ore_attese += HOURS_EXPECTED.get(current.weekday(), 8)
            current += timedelta(days=1)

        self.ui.lbl_summary.config(
            text=(f"{len(df_s):,} records  │  Total hours: {total:.1f}  │  "
                  f"Expected: {ore_attese:.0f}h  │  "
                  f"Period: {d_from} → {d_to}"))

    # ── Sort ──────────────────────────────────────────────────────────────────

    def _sort(self, col):
        items = [(self.ui.tree.set(c, col), c)
                 for c in self.ui.tree.get_children("")]
        asc = not getattr(self, f"_asc_{col}", True)
        setattr(self, f"_asc_{col}", asc)
        items.sort(key=lambda x: x[0], reverse=not asc)
        for idx, (_, item) in enumerate(items):
            self.ui.tree.move(item, "", idx)

    # ── Export ────────────────────────────────────────────────────────────────

    def _export(self):
        df = self._last_df
        if df is None or df.empty:
            messagebox.showwarning("Export", "No data to export.")
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
            title="Save results")
        if not path:
            return
        try:
            import openpyxl
            from openpyxl.styles import (
                PatternFill, Font, Alignment, Border, Side
            )
            from openpyxl.utils import get_column_letter
            from openpyxl.worksheet.table import Table, TableStyleInfo

            # Colonne da esportare con etichette
            col_map = {
                "DescrResource": "Resource",
                "CDCOM":         "Job",
                "DescrFase":     "Phase",
                "DescrShort":    "Description",
                "WorkNote":      "Notes",
                "WorkDte":       "Date",
                "CreatedOn":     "Created On",
                "WorkHrs":       "Hours",
            }
            # Prendi solo le colonne esistenti
            cols_avail = [c for c in col_map if c in df.columns]
            df_exp = df[cols_avail].copy()
            df_exp.columns = [col_map[c] for c in cols_avail]

            # Converti date in stringa
            for c in ("Date", "Created On"):
                if c in df_exp.columns:
                    df_exp[c] = df_exp[c].astype(str)

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Hours"

            # Scrivi intestazioni
            headers = list(df_exp.columns)
            for ci, h in enumerate(headers, 1):
                ws.cell(row=1, column=ci, value=h)

            # Scrivi dati
            for ri, row in enumerate(df_exp.itertuples(index=False), 2):
                for ci, val in enumerate(row, 1):
                    ws.cell(row=ri, column=ci, value=val)

            n_rows = len(df_exp)
            n_cols = len(headers)

            # Crea tabella Excel con filtri automatici
            tbl_ref = f"A1:{get_column_letter(n_cols)}{n_rows + 1}"
            tbl = Table(displayName="HoursData", ref=tbl_ref)
            tbl.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False
            )
            ws.add_table(tbl)

            # Riga totale sotto la tabella
            total_row = n_rows + 2
            hrs_col = headers.index("Hours") + 1 if "Hours" in headers else None

            # Stili
            accent_fill  = PatternFill("solid", fgColor="9a3412")
            total_fill   = PatternFill("solid", fgColor="dcfce7")
            white_font   = Font(name="Segoe UI", bold=True, color="FFFFFF", size=10)
            total_font   = Font(name="Segoe UI", bold=True, color="166534", size=10)
            normal_font  = Font(name="Segoe UI", size=10)
            center_align = Alignment(horizontal="center", vertical="center")
            left_align   = Alignment(horizontal="left",   vertical="center")
            thin = Side(style="thin", color="e2e5ea")
            thin_border  = Border(top=thin, bottom=thin, left=thin, right=thin)

            # Formatta intestazioni
            for ci in range(1, n_cols + 1):
                cell = ws.cell(row=1, column=ci)
                cell.fill   = accent_fill
                cell.font   = white_font
                cell.alignment = center_align

            # Formatta dati
            for ri in range(2, n_rows + 2):
                for ci in range(1, n_cols + 1):
                    cell = ws.cell(row=ri, column=ci)
                    cell.font      = normal_font
                    cell.alignment = center_align if headers[ci-1] in ("Date", "Created On", "Hours") else left_align

            # Riga totale
            if hrs_col:
                ws.cell(row=total_row, column=hrs_col - 1, value="TOTAL HOURS →").font = total_font
                ws.cell(row=total_row, column=hrs_col - 1).fill = total_fill
                ws.cell(row=total_row, column=hrs_col - 1).alignment = Alignment(horizontal="right")
                # SUBTOTAL(9,...) somma solo le righe visibili dopo i filtri
                hrs_col_letter = get_column_letter(hrs_col)
                formula = f"=SUBTOTAL(9,{hrs_col_letter}2:{hrs_col_letter}{n_rows + 1})"
                ws.cell(row=total_row, column=hrs_col, value=formula).font = total_font
                ws.cell(row=total_row, column=hrs_col).fill = total_fill
                ws.cell(row=total_row, column=hrs_col).alignment = center_align

            # Adatta larghezza colonne
            col_widths = {
                "Resource": 22, "Job": 16, "Phase": 22,
                "Description": 22, "Notes": 35,
                "Date": 12, "Created On": 12, "Hours": 10,
            }
            for ci, h in enumerate(headers, 1):
                ws.column_dimensions[get_column_letter(ci)].width = col_widths.get(h, 15)

            # Altezza righe
            ws.row_dimensions[1].height = 20
            for ri in range(2, n_rows + 2):
                ws.row_dimensions[ri].height = 18

            # Freezing intestazioni
            ws.freeze_panes = "A2"

            wb.save(path)
            messagebox.showinfo("Export completed", f"Saved:\n{path}")

        except Exception as e:
            messagebox.showerror("Export error", str(e))

    # ── Date picker ───────────────────────────────────────────────────────────

    def _pick_date(self, entry_widget):
        try:
            from tkcalendar import Calendar
        except ImportError:
            messagebox.showwarning("Missing",
                "Install tkcalendar:\n  pip install tkcalendar")
            return

        current = parse_date(entry_widget.get().strip()) or date.today()
        top = tk.Toplevel(self.ui.root)
        top.title("Select Date")
        top.resizable(False, False)
        top.grab_set()
        top.configure(bg=BG_CARD)
        set_app_icon(top)

        cal = Calendar(top, selectmode="day",
                       year=current.year, month=current.month, day=current.day,
                       date_pattern="yyyy-mm-dd",
                       background=ACCENT, foreground=TEXT_WHITE,
                       headersbackground=ACCENT_MID, headersforeground=TEXT_WHITE,
                       selectbackground=ACCENT_LIGHT, selectforeground=TEXT,
                       normalbackground=BG_CARD, normalforeground=TEXT,
                       weekendbackground=BG_CARD, weekendforeground=TEXT_DIM,
                       font=("Segoe UI", 10))
        cal.pack(padx=10, pady=10)

        def confirm():
            entry_widget.delete(0, tk.END)
            entry_widget.insert(0, cal.get_date())
            top.destroy()

        tk.Button(top, text="Confirm", command=confirm,
                  font=("Segoe UI", 10), bg=ACCENT, fg=TEXT_WHITE,
                  relief="flat", cursor="hand2", padx=16, pady=6).pack(pady=(0, 10))

    # ── Update check ──────────────────────────────────────────────────────────

    def _check_update_available(self):
        threading.Thread(target=self._check_thread, daemon=True).start()

    def _check_thread(self):
        try:
            xlsx_mtime = os.path.getmtime(self.xlsx_path)
            xlsx_dt    = datetime.fromtimestamp(xlsx_mtime)
            df, exported_at = json_to_df(JSON_FILE)
            json_dt = datetime.strptime(exported_at, "%Y-%m-%d %H:%M:%S")
            if xlsx_dt > json_dt:
                hours = int((xlsx_dt - json_dt).total_seconds() // 3600)
                msg   = f"⚠  Updated data available on server (Excel modified {hours}h ago)"
                self.ui.root.after(0, lambda: self._show_banner(msg))
        except Exception:
            pass

    def _show_banner(self, msg):
        self.ui.lbl_banner.config(text=msg)
        self.ui.frm_banner.pack(fill="x", before=self.ui.lbl_summary)

    # ── Run ───────────────────────────────────────────────────────────────────

    def run(self):
        self.ui.root.mainloop()


# ── Entry point ───────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="Engineering Hours Viewer")
    parser.add_argument("--xlsx", default=XLSX_FILE)
    args = parser.parse_args()
    HoursViewer(xlsx_path=args.xlsx).run()


if __name__ == "__main__":
    main()